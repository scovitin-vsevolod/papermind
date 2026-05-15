"""Smoke tests for the markitdown loader.

Fixtures are generated on the fly so no binary files are checked into
the repo. python-docx and fpdf2 are dev-only dependencies.
"""

from __future__ import annotations

import io

import pytest
from docx import Document as DocxDocument
from fpdf import FPDF

from app.loaders.markitdown_loader import LoaderError, to_markdown

# ─── fixture builders ────────────────────────────────────────────────────────


def _build_docx(text: str) -> bytes:
    doc = DocxDocument()
    for paragraph in text.split("\n"):
        doc.add_paragraph(paragraph)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_pdf(text: str) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for line in text.split("\n"):
        pdf.cell(0, 10, line, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


# ─── tests ───────────────────────────────────────────────────────────────────


def test_txt_loader():
    text = "Hello world.\nSecond line of plain text."
    result = to_markdown(text.encode("utf-8"), "sample.txt")
    assert "Hello world" in result
    assert "Second line" in result


def test_md_loader_preserves_headings():
    md = "# Title\n\nBody paragraph with **bold**."
    result = to_markdown(md.encode("utf-8"), "sample.md")
    assert "Title" in result
    assert "bold" in result


def test_docx_loader():
    content = _build_docx("Hello DOCX.\nThis is a second paragraph.")
    result = to_markdown(content, "sample.docx")
    assert "Hello DOCX" in result
    assert "second paragraph" in result


def test_pdf_loader():
    content = _build_pdf("Hello PDF.\nLine two of the PDF.")
    result = to_markdown(content, "sample.pdf")
    assert "Hello PDF" in result


def test_missing_extension_raises():
    with pytest.raises(LoaderError, match="missing extension"):
        to_markdown(b"whatever", "noext")


def test_empty_content_raises():
    with pytest.raises(LoaderError, match="empty text"):
        to_markdown(b"", "empty.txt")
